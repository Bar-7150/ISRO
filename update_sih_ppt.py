from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

FILE = r"c:\Users\sunet\Documents\ISRO\NMIET_SIH_2026_PPT_Template.pptx"

prs = Presentation(FILE)

team_name = "ISRO Reliability Analytics Team"

slide_1 = prs.slides[0]
for shape in slide_1.shapes:
    if hasattr(shape, "text"):
        if shape.name == "Title 7":
            shape.text = "SMART INDIA HACKATHON 2026"
        elif shape.name == "Subtitle 3":
            shape.text = "TITLE PAGE"
        elif shape.name == "TextBox 9":
            shape.text = "PS ID – ISRO-2026 | Problem Statement Title – AI-Powered Burn-In Screening for High-Reliability Semiconductor Components | Theme – AI / Quality & Reliability | PS Category – Software | Team ID – ISRO-AI-QA | Team Name – ISRO Reliability Analytics Team"
        elif "Oval" in shape.name or shape.text.strip() == "Your Team Name":
            shape.text = team_name

# Slide 2: Problem / Idea
slide_2 = prs.slides[1]
for shape in slide_2.shapes:
    if hasattr(shape, "text"):
        if "Title" in shape.name:
            shape.text = "IDEA TITLE"
        elif shape.name == "TextBox 8":
            shape.text = "Proposed Solution: AI-driven burn-in screening system for semiconductor reliability\n\n• Detects early lot-level anomalies and latent defects before they escape into high-risk deployment\n• Combines contextual anomaly detection with predictive drift forecasting to flag unstable devices at 24h/96h stages\n• Provides explainable QA outputs using SHAP and decision rules for engineering sign-off\n\nWhy it matters: conventional static limits miss defects that appear normal early but drift catastrophically later in burn-in."
        elif "Oval" in shape.name or shape.text.strip() == "Your Team Name":
            shape.text = team_name

# Slide 3: Technical Approach
slide_3 = prs.slides[2]
for shape in slide_3.shapes:
    if hasattr(shape, "text"):
        if "Title" in shape.name:
            shape.text = "TECHNICAL APPROACH"
        elif shape.name == "TextBox 8":
            shape.text = "Technologies used: Python, scikit-learn, XGBoost, SHAP, Matplotlib, Seaborn, pandas, NumPy\n\nMethodology:\n• Feature engineering: Δ0–24h drift, % drift, lot-normalized Z-score, stress interaction signals\n• Module A: Mahalanobis + Isolation Forest + LOF ensemble for contextual anomaly detection\n• Module B: Gradient boosting with MAE + Bayesian uncertainty for 168h degradation prediction\n• Decision layer: cost-sensitive Fβ thresholding and audit-ready QA disposition report"
        elif "Oval" in shape.name or shape.text.strip() == "Your Team Name":
            shape.text = team_name

# Slide 4: Feasibility & Viability
slide_4 = prs.slides[3]
for shape in slide_4.shapes:
    if hasattr(shape, "text"):
        if "Title" in shape.name:
            shape.text = "FEASIBILITY AND VIABILITY"
        elif shape.name == "TextBox 8":
            shape.text = "Feasibility:\n• Uses standard industrial data pipelines and open-source ML libraries\n• Works on burn-in production logs and wafer-level sensor data\n• No special hardware required; deployment is software-based with manufacturing integration\n\nChallenges and mitigation:\n• Class imbalance → cost-sensitive scoring and recall-focused thresholding\n• Lot-to-lot drift → contextual normalization and drift features\n• Explainability need → SHAP + rule-based QA summaries"
        elif "Oval" in shape.name or shape.text.strip() == "Your Team Name":
            shape.text = team_name

# Slide 5: Impact and Benefits
slide_5 = prs.slides[4]
for shape in slide_5.shapes:
    if hasattr(shape, "text"):
        if "Title" in shape.name:
            shape.text = "IMPACT AND BENEFITS"
        elif shape.name == "TextBox 8":
            shape.text = "Target audience:\n• Semiconductor fabs, reliability labs, aerospace component QA teams, defense electronics producers\n\nBenefits:\n• Reduces escapes and safety-critical field failures\n• Improves screening yield without sacrificing quality\n• Enables early abort decisions for unstable lots\n• Supports transparent engineering review and auditability"
        elif "Oval" in shape.name or shape.text.strip() == "Your Team Name":
            shape.text = team_name

# Slide 6: Research and References
slide_6 = prs.slides[5]
for shape in slide_6.shapes:
    if hasattr(shape, "text"):
        if "Title" in shape.name:
            shape.text = "RESEARCH AND REFERENCES"
        elif shape.name == "TextBox 8":
            shape.text = "Reference sources and benchmark mapping:\n• UCI SECOM semiconductor defect dataset for contextual anomaly screening\n• NASA C-MAPSS turbofan degradation dataset for time-to-failure drift prediction\n• Reliability engineering literature on burn-in, latent defect detection, and drift acceleration\n• Project implementation: end-to-end ML pipeline for early anomaly detection and late-stage degradation forecasting"
        elif "Oval" in shape.name or shape.text.strip() == "Your Team Name":
            shape.text = team_name

# Add chart visuals to slides 2-5 for a more demo-ready SIH deck
try:
    from pathlib import Path
    fig_dir = Path(r"c:\Users\sunet\Documents\ISRO\reports\figures")
    image_map = {
        1: fig_dir / "kaggle_secom_roc_pr_curves.png",
        2: fig_dir / "kaggle_secom_confusion_matrices.png",
        3: fig_dir / "kaggle_cmapss_loss_curve.png",
        4: fig_dir / "kaggle_cmapss_residual_diagnostics.png",
    }
    for slide_index, image_name in image_map.items():
        if image_name.exists():
            slide = prs.slides[slide_index]
            left = 7.5
            top = 1.9
            width = 4.7
            height = 3.6
            slide.shapes.add_picture(str(image_name), left, top, width, height)
except Exception:
    pass

prs.save(FILE)
print(f"Updated SIH presentation saved to: {FILE}")
