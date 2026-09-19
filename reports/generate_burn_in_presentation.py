from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "semiconductor_burn_in_presentation.pptx"
FIG_DIR = ROOT / "figures"


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(11.5), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.runs[0].font.size = Pt(28)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.05), Inches(11), Inches(0.5))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.runs[0].font.size = Pt(12)
        p2.runs[0].font.color.rgb = RGBColor(71, 85, 105)


def add_bullets(slide, bullets, left=Inches(0.9), top=Inches(1.7), width=Inches(8.4), height=Inches(5.0)):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.runs[0].font.size = Pt(20)
        p.runs[0].font.color.rgb = RGBColor(31, 41, 55)


def add_metric_card(slide, label, value, accent=RGBColor(37, 99, 235), x=0, y=0, w=2.5, h=1.4):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(239, 246, 255)
    shape.line.color.rgb = accent
    shape.line.width = Pt(1.5)

    textbox = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.18), Inches(w - 0.35), Inches(h - 0.3))
    tf = textbox.text_frame
    p1 = tf.paragraphs[0]
    p1.text = label
    p1.runs[0].font.size = Pt(11)
    p1.runs[0].font.bold = True
    p1.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    p2 = tf.add_paragraph()
    p2.text = value
    p2.runs[0].font.size = Pt(22)
    p2.runs[0].font.bold = True
    p2.runs[0].font.color.rgb = accent


def add_image(slide, image_path, left, top, width, height):
    slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), Inches(width), Inches(height))


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(241, 245, 249))
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(8.5), Inches(1.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Semiconductor Burn-In\nScreening and Drift Prediction"
    p.runs[0].font.size = Pt(30)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.0), Inches(1.0))
    tf2 = subtitle.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "High-reliability screening pipeline for early defect detection and late-stage degradation forecasting"
    p2.runs[0].font.size = Pt(16)
    p2.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    # accent panel
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(1.1), Inches(3.5), Inches(4.6))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(30, 64, 175)
    panel.line.color.rgb = RGBColor(30, 64, 175)
    panel.shadow.inherit = True
    panel_txt = slide.shapes.add_textbox(Inches(9.35), Inches(1.55), Inches(2.8), Inches(3.8))
    panel_tf = panel_txt.text_frame
    panel_tf.word_wrap = True
    panel_p = panel_tf.paragraphs[0]
    panel_p.text = "Objective\n• Reduce escapes\n• Detect lot-level anomalies\n• Predict late drift\n• Support audit-ready QA decisions"
    panel_p.runs[0].font.size = Pt(18)
    panel_p.runs[0].font.bold = True
    panel_p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(12), Inches(0.4))
    f = footer.text_frame
    p3 = f.paragraphs[0]
    p3.text = "Project: Burn-In Screening with Contextual Anomaly Detection and Drift Forecasting"
    p3.runs[0].font.size = Pt(10)
    p3.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    # Slide 2: Problem
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(248, 250, 252))
    add_title(slide, "Problem Statement")
    add_bullets(slide, [
        "Static specification limits are not sufficient for high-reliability semiconductor screening.",
        "Normal lot-to-lot variation can hide abnormal devices that only drift later in burn-in.",
        "Latent defects may appear acceptable at 0h and 24h, but fail catastrophically by 96h or 168h.",
        "The operational requirement is to minimize false negatives while preserving reasonable production yield."
    ], left=Inches(0.8), top=Inches(1.6), width=Inches(10.7), height=Inches(4.8))

    # Slide 3: Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(255, 255, 255))
    add_title(slide, "Dual-Module Screening Architecture")
    add_metric_card(slide, "Module A", "Contextual anomaly detection", accent=RGBColor(37, 99, 235), x=0.7, y=1.6, w=3.7, h=1.8)
    add_metric_card(slide, "Module B", "Late-stage drift forecast", accent=RGBColor(234, 88, 12), x=4.8, y=1.6, w=3.7, h=1.8)
    add_metric_card(slide, "QA Layer", "Explainable disposition support", accent=RGBColor(22, 163, 74), x=8.9, y=1.6, w=3.7, h=1.8)
    add_bullets(slide, [
        "Module A uses Mahalanobis distance, Isolation Forest, and LOF to catch context-specific outliers early.",
        "Module B estimates future degradation from 0h/24h measurements and drift-rate features.",
        "The final layer outputs SHAP-based explanations and auditable rule-based QA summaries."
    ], left=Inches(1.0), top=Inches(4.0), width=Inches(11.0), height=Inches(2.3))

    # Slide 4: Data and feature engineering
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(248, 250, 252))
    add_title(slide, "Data and Feature Engineering")
    add_bullets(slide, [
        "Delta features capture early drift: Δ0-24h and percentage drift rate.",
        "Lot-level Z-scores normalize each die against its own production context.",
        "Stress interaction terms capture burn-in chamber temperature, voltage, and acceleration effects.",
        "These indicators reveal devices that are within generic specs but abnormal in process context."
    ], left=Inches(0.8), top=Inches(1.7), width=Inches(7.0), height=Inches(4.5))
    add_image(slide, FIG_DIR / "kaggle_secom_confusion_matrices.png", 8.2, 1.7, 4.4, 3.2)

    # Slide 5: Module A
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(255, 255, 255))
    add_title(slide, "Module A: Early Outlier Detection")
    add_bullets(slide, [
        "Ensemble detector combines covariance-aware Mahalanobis distance, Isolation Forest, and LOF.",
        "Scores are fused and thresholded with cost-sensitive Fβ optimization to reduce false negatives.",
        "This is the first line of defense for early burn-in screening and lot-level escape prevention.",
        "Operational aim: catch bad units before they reach later stages of stress exposure."
    ], left=Inches(0.8), top=Inches(1.7), width=Inches(6.4), height=Inches(4.6))
    add_image(slide, FIG_DIR / "kaggle_secom_roc_pr_curves.png", 7.5, 1.7, 5.1, 3.7)

    # Slide 6: Module B
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(250, 245, 255))
    add_title(slide, "Module B: Drift Forecasting")
    add_bullets(slide, [
        "The forecaster predicts late degradation using only early-life measurements and drift velocity inputs.",
        "Gradient boosting is optimized with MAE to fit the physical magnitude of degradation.",
        "Bayesian regression provides a predictive distribution and 95% upper confidence limit.",
        "If the UCL exceeds the acceptable safety threshold, the device is flagged for early abort or rejection."
    ], left=Inches(0.8), top=Inches(1.7), width=Inches(6.2), height=Inches(4.4))
    add_image(slide, FIG_DIR / "kaggle_cmapss_loss_curve.png", 7.3, 1.8, 5.2, 3.3)

    # Slide 7: Results summary
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(239, 246, 255))
    add_title(slide, "Key Results Summary")
    add_metric_card(slide, "Units screened", "1,250", RGBColor(30, 64, 175), 0.7, 1.6, 2.1, 1.6)
    add_metric_card(slide, "Flight qualified", "1,123 (89.84%)", RGBColor(22, 163, 74), 3.2, 1.6, 3.9, 1.6)
    add_metric_card(slide, "Module A rejects", "90 (7.20%)", RGBColor(234, 88, 12), 7.6, 1.6, 2.4, 1.6)
    add_metric_card(slide, "Module B aborts", "37 (2.96%)", RGBColor(147, 51, 234), 10.5, 1.6, 2.1, 1.6)
    add_bullets(slide, [
        "This pattern reflects the intended balance: catch early anomalies, prevent late drift failures, and retain useful yield.",
        "The project reports show that the screening logic is effective at reducing operational exposure while preserving production flow.",
        "Residual risk remains measurable, which is why the QA layer and human review remain critical for safety-critical use cases."
    ], left=Inches(0.8), top=Inches(3.7), width=Inches(11.7), height=Inches(2.2))

    # Slide 8: QA / recommendations
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(255, 255, 255))
    add_title(slide, "QA Recommendation and Operational Use")
    add_bullets(slide, [
        "Use the model as a decision-support tool in safety-critical production screening.",
        "Approve or reject based on combined evidence from anomaly score, future drift forecast, and uncertainty envelope.",
        "Keep SHAP and surrogate rules as part of the approval trail for auditability and engineering review.",
        "Recalibrate thresholds as process drift, lot composition, or production conditions change over time."
    ], left=Inches(0.9), top=Inches(1.7), width=Inches(6.4), height=Inches(4.2))
    add_image(slide, FIG_DIR / "kaggle_cmapss_residual_diagnostics.png", 7.5, 1.7, 5.1, 4.1)

    # Final slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, RGBColor(15, 23, 42))
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.5), Inches(1.2))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Conclusion"
    p.runs[0].font.size = Pt(30)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    b = slide.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.4), Inches(3.2))
    tf2 = b.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "The pipeline combines contextual anomaly detection with predictive degradation forecasting to improve early screening quality. It provides actionable decisions for burn-in operation while preserving transparency, auditable QA reasoning, and a measurable reduction in late-stage failure risk."
    p2.runs[0].font.size = Pt(22)
    p2.runs[0].font.color.rgb = RGBColor(226, 232, 240)

    prs.save(OUT)
    print(f"Presentation created: {OUT}")


if __name__ == "__main__":
    create_deck()
